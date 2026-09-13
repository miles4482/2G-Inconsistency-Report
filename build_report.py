#!/usr/bin/env python3
"""Rebuild comparison reports with case-insensitive Huawei parameter descriptions."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pyxlsb import open_workbook

JSON_PATH = Path("/workspace/2G_Parameter_Change_Data.json")
PARAM_LIST = "/workspace/GSM_Parameter List_Huawei_v1.0.xlsb"
OUT_XLSX = Path("/workspace/2G_Parameter_Change_Summary.xlsx")
OUT_CSV = Path("/workspace/2G_Changed_Parameters_Ranked.csv")
OUT_MD = Path("/workspace/2G_Parameter_Change_Report.md")

CDR_STRONG = re.compile(
    r"\b(call drop|drop rate|handover|srvcc|vamos|tight bcch|tch/h|half.rate|"
    r"voice quality|mos|bsic|ncc|bcc|max ta|fast return|csfb)\b",
    re.I,
)


def load_catalog():
    cat = {}
    with open_workbook(PARAM_LIST) as wb:
        with wb.get_sheet(1) as sh:
            it = sh.rows()
            header = [c.v for c in next(it)]
            idx = {str(n): i for i, n in enumerate(header) if n}

            def get(vals, col):
                i = idx.get(col)
                if i is None or i >= len(vals) or vals[i] is None:
                    return ""
                return str(vals[i]).strip()

            for row in it:
                vals = [c.v for c in row]
                mo = get(vals, "MO").upper()
                pid = get(vals, "Parameter ID").upper()
                if not mo or not pid or (mo, pid) in cat:
                    continue
                cat[(mo, pid)] = {
                    "name": get(vals, "Parameter Name"),
                    "meaning": get(vals, "Meaning"),
                    "impact": get(vals, "Impact on Radio Network Performance"),
                    "recommended": get(vals, "Recommended Value"),
                    "default": get(vals, "Default Value"),
                    "range": get(vals, "GUI Value Range"),
                }
    return cat


def classify(r):
    """Return (change_type, kpi_focus)."""
    n = r["n_changed"]
    n_trans = r["n_trans"]
    trans = r["transitions"]
    pid = r["pid"].upper()
    mo = r["mo"].upper()
    blob = " ".join([pid, r.get("pname", ""), r.get("meaning", ""), r.get("impact", "")])

    if n >= 20000:
        ctype = "Network-wide bulk change"
    elif n >= 500 and n_trans <= 5:
        ctype = "Cluster bulk change"
    elif n_trans >= 50:
        ctype = "Per-cell / SON-like drift"
    else:
        ctype = "Limited / targeted change"

    kpi = "Other"
    if pid in {"NCC", "BCC", "LAC"}:
        kpi = "HO / identity (can cause HO fail & drop)"
    elif "VAMOS" in pid or mo == "GCELLVAMOS":
        kpi = "TCH CDR / MOS (VAMOS quality)"
    elif pid in {"TCHBUSYTHRES", "AMRTCHHPRIORLOAD"}:
        kpi = "TCH CDR / MOS (more TCHH)"
    elif "TIGHTBCCH" in pid:
        kpi = "TCH CDR / CSSR (tight BCCH reuse)"
    elif pid == "IMMREJWAITINDTIMER" or "INTACESCONG" in pid:
        kpi = "CSSR / SDCCH (access delay); possible SDCCH drop"
    elif "LTESAI" in pid or "SRVCC" in pid:
        kpi = "SRVCC HO / CS continuity (VoLTE→GSM)"
    elif "FASTRETURN" in pid:
        kpi = "CSFB/SRVCC return & ping-pong"
    elif pid in {"DYNOPENTRXPOWER", "ENERGYCONSRVPREFSW", "PDCHTRXDYNSHUTSW", "CONCENINTELSHUTDOWNSW"}:
        kpi = "Energy saving (coverage/drop if ON)"
    elif pid == "MAXTA":
        kpi = "TCH CDR (timing / coverage)"
    elif CDR_STRONG.search(blob):
        kpi = "Radio KPI related"
    return ctype, kpi


def enrich(data, cat):
    for r in data["summary"]:
        meta = cat.get((r["mo"].upper(), r["pid"].upper()), {})
        if meta.get("meaning"):
            r["meaning"] = meta["meaning"]
        if meta.get("impact"):
            r["impact"] = meta["impact"]
        if meta.get("recommended"):
            r["recommended"] = meta["recommended"]
        if meta.get("default"):
            r["default"] = meta["default"]
        if meta.get("range"):
            r["range"] = meta["range"]
        if meta.get("name") and (not r.get("pname") or r["pname"] == r["pid"]):
            r["pname"] = meta["name"]
        r["change_type"], r["kpi"] = classify(r)
        r["cdr"] = r["kpi"] != "Other"
    data["summary"].sort(key=lambda x: (-x["n_changed"], x["mo"], x["pid"]))
    return data


def write_csv(rows):
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "Rank",
                "MO",
                "Parameter ID",
                "Parameter Name",
                "Cells Changed",
                "Pct Common Cells",
                "Common Cells",
                "Distinct Transitions",
                "Change Type",
                "Likely KPI Impact",
                "Top Value Changes (Pre → Post)",
                "Default Value",
                "Recommended Value",
                "Meaning",
                "Impact on Radio Network Performance",
            ]
        )
        for i, r in enumerate(rows, 1):
            w.writerow(
                [
                    i,
                    r["mo"],
                    r["pid"],
                    r["pname"],
                    r["n_changed"],
                    r["pct"],
                    r["n_common"],
                    r["n_trans"],
                    r["change_type"],
                    r["kpi"],
                    r["transitions"],
                    r.get("default", ""),
                    r.get("recommended", ""),
                    r.get("meaning", ""),
                    r.get("impact", ""),
                ]
            )


def write_xlsx(data):
    rows = data["summary"]
    stats = data["sheet_stats"]
    new_cells = data["new_cells"]
    del_cells = data["del_cells"]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    alt_fill = PatternFill("solid", fgColor="D6EAF8")
    top_fill = PatternFill("solid", fgColor="F5B7B1")
    bulk_fill = PatternFill("solid", fgColor="FDEBD0")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, ncols):
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        for col in range(1, ncols + 1):
            cell = ws.cell(1, col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def autosize(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def paint(ws, ridx, ncols, fill):
        for col in range(1, ncols + 1):
            cell = ws.cell(ridx, col)
            cell.border = thin
            cell.alignment = wrap
            if fill:
                cell.fill = fill

    headers = [
        "Rank",
        "MO",
        "MO Description",
        "Parameter ID",
        "Parameter Name",
        "Cells Changed",
        "% of Common Cells",
        "Common Cells",
        "Distinct Transitions",
        "Change Type",
        "Likely KPI Impact",
        "Top Value Changes (Pre → Post)",
        "Default Value",
        "Recommended Value",
        "Meaning",
        "Impact on Radio Network Performance",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Changed Parameters (Ranked)"
    ws.append(headers)
    style_header(ws, len(headers))
    for i, r in enumerate(rows, 1):
        ws.append(
            [
                i,
                r["mo"],
                r.get("mo_desc", ""),
                r["pid"],
                r["pname"],
                r["n_changed"],
                r["pct"],
                r["n_common"],
                r["n_trans"],
                r["change_type"],
                r["kpi"],
                r["transitions"],
                r.get("default", ""),
                r.get("recommended", ""),
                r.get("meaning", ""),
                r.get("impact", ""),
            ]
        )
        fill = top_fill if i <= 15 else (bulk_fill if r["n_changed"] >= 20000 else (alt_fill if i % 2 == 0 else None))
        paint(ws, i + 1, len(headers), fill)
        ws.row_dimensions[i + 1].height = 46 if i <= 30 else 32
    autosize(ws, [8, 20, 34, 28, 42, 16, 18, 14, 16, 26, 36, 72, 18, 20, 70, 55])

    ws2 = wb.create_sheet("Network-wide Bulk Changes")
    ws2.append(headers)
    style_header(ws2, len(headers))
    bulk = [r for r in rows if r["n_changed"] >= 5000]
    for i, r in enumerate(bulk, 1):
        ws2.append(
            [
                i,
                r["mo"],
                r.get("mo_desc", ""),
                r["pid"],
                r["pname"],
                r["n_changed"],
                r["pct"],
                r["n_common"],
                r["n_trans"],
                r["change_type"],
                r["kpi"],
                r["transitions"],
                r.get("default", ""),
                r.get("recommended", ""),
                r.get("meaning", ""),
                r.get("impact", ""),
            ]
        )
        paint(ws2, i + 1, len(headers), top_fill if i <= 9 else alt_fill)
        ws2.row_dimensions[i + 1].height = 55
    autosize(ws2, [8, 20, 34, 28, 42, 16, 18, 14, 16, 26, 36, 72, 18, 20, 70, 55])

    ws3 = wb.create_sheet("CDR Suspects")
    ws3.append(headers)
    style_header(ws3, len(headers))
    suspects = [
        r
        for r in rows
        if any(
            k in r["kpi"]
            for k in ("TCH CDR", "SRVCC", "HO / identity", "tight BCCH", "CSFB")
        )
    ]
    for i, r in enumerate(suspects, 1):
        ws3.append(
            [
                i,
                r["mo"],
                r.get("mo_desc", ""),
                r["pid"],
                r["pname"],
                r["n_changed"],
                r["pct"],
                r["n_common"],
                r["n_trans"],
                r["change_type"],
                r["kpi"],
                r["transitions"],
                r.get("default", ""),
                r.get("recommended", ""),
                r.get("meaning", ""),
                r.get("impact", ""),
            ]
        )
        paint(ws3, i + 1, len(headers), top_fill if i <= 10 else (alt_fill if i % 2 == 0 else None))
        ws3.row_dimensions[i + 1].height = 48
    autosize(ws3, [8, 20, 34, 28, 42, 16, 18, 14, 16, 26, 36, 72, 18, 20, 70, 55])

    ws4 = wb.create_sheet("MO Comparison Stats")
    h4 = [
        "MO",
        "MO Description",
        "Pre Cells",
        "Post Cells",
        "Common Cells",
        "New in Post",
        "Removed since Pre",
        "Params Compared",
        "Params Changed",
        "Total Cell-Param Changes",
        "Params only in Pre",
        "Params only in Post",
    ]
    ws4.append(h4)
    style_header(ws4, len(h4))
    ordered = sorted(stats, key=lambda x: -x["n_total_changes"])
    for i, s in enumerate(ordered, 1):
        ws4.append(
            [
                s["mo"],
                s["mo_desc"],
                s["n_pre"],
                s["n_post"],
                s["n_common"],
                s["n_new"],
                s["n_del"],
                s["n_params"],
                s["n_params_changed"],
                s["n_total_changes"],
                s["only_pre"],
                s["only_post"],
            ]
        )
        paint(ws4, i + 1, len(h4), alt_fill if i % 2 == 0 else None)
    autosize(ws4, [24, 42, 12, 12, 14, 14, 18, 16, 16, 24, 28, 28])

    g = next((x for x in stats if x["mo"] == "GCELL"), None)
    ws5 = wb.create_sheet("Inventory")
    ws5.append(["Item", "Value"])
    style_header(ws5, 2)
    inv = [
        ("Pre dump", "2G Cell Parameters Data_26April26_Pre"),
        ("Post dump", "2G Cell Parameters Data_13Sep26_Post"),
        ("Pre cells", g["n_pre"] if g else ""),
        ("Post cells", g["n_post"] if g else ""),
        ("Common cells compared", g["n_common"] if g else ""),
        ("New cells in Post", data["n_new_cells"]),
        ("Cells removed since Pre", data["n_del_cells"]),
        ("Parameters with any value change", len(rows)),
        ("Network-wide bulk parameters (>=5k cells)", len(bulk)),
    ]
    for i, (a, b) in enumerate(inv, 1):
        ws5.append([a, b])
        paint(ws5, i + 1, 2, alt_fill if i % 2 == 0 else None)
    autosize(ws5, [52, 48])

    ws6 = wb.create_sheet("New Cells (Post only)")
    ws6.append(["BSCName", "CELLNAME"])
    style_header(ws6, 2)
    for i, k in enumerate(new_cells, 1):
        ws6.append([k[0], k[1]])
        paint(ws6, i + 1, 2, alt_fill if i % 2 == 0 else None)
    autosize(ws6, [18, 24])

    ws7 = wb.create_sheet("Removed Cells (Pre only)")
    ws7.append(["BSCName", "CELLNAME"])
    style_header(ws7, 2)
    for i, k in enumerate(del_cells, 1):
        ws7.append([k[0], k[1]])
        paint(ws7, i + 1, 2, alt_fill if i % 2 == 0 else None)
    autosize(ws7, [18, 24])

    wb.save(OUT_XLSX)


def clip(s, n=450):
    s = (s or "").replace("\n", " ").strip()
    return s if len(s) <= n else s[: n - 3] + "..."


def write_md(data):
    rows = data["summary"]
    stats = data["sheet_stats"]
    g = next((x for x in stats if x["mo"] == "GCELL"), None)
    bulk = [r for r in rows if r["n_changed"] >= 5000]

    lines = []
    a = lines.append
    a("# 2G Cell Parameter Comparison — Pre vs Post")
    a("")
    a("| Item | Value |")
    a("|---|---|")
    a("| Pre dump | 26 April 2026 (`2G Cell Parameters Data_26April26_Pre`) |")
    a("| Post dump | 13 September 2026 (`2G Cell Parameters Data_13Sep26_Post`) |")
    a("| Descriptions | `GSM_Parameter List_Huawei_v1.0` |")
    if g:
        a(f"| Pre cells | {g['n_pre']:,} |")
        a(f"| Post cells | {g['n_post']:,} |")
        a(f"| Common cells compared | {g['n_common']:,} |")
    a(f"| New cells in Post | {data['n_new_cells']:,} |")
    a(f"| Cells removed since Pre | {data['n_del_cells']:,} |")
    a(f"| Parameters with any change | {len(rows)} |")
    a("")
    a("Identity columns (BSC/BTS/Cell name/index) were excluded. Rank is by **number of common cells whose value changed**.")
    a("")
    a("Classic TCH drop parameters (handover hysteresis/thresholds, RLT/radio-link timers, power control, RXQUAL HO) **did not change**. The CDR rise lines up with a small set of **network-wide bulk campaigns** plus VAMOS / half-rate / identity changes.")
    a("")
    a("## Highest-priority findings for sudden CDR increase")
    a("")
    a("### 1. Intelligent Access Congestion Control enabled on almost the whole network")
    a("")
    a("On **24,541 cells (98.9%)**:")
    a("")
    a("- `INTACESCONGCTRLSW`: **OFF → ON** (Huawei default/recommended = OFF)")
    a("- `INTACESCONGCTRLTHRES`: empty → **5** (recommended **20** — much more sensitive)")
    a("- `INTACESCONGCTRLTIMER1`: empty → **1** (recommended **2** — T3122 increases faster)")
    a("- `INTACESCONGCTRLTIMER2`: empty → **100** (recommended **5** — inflated T3122 stays high far longer)")
    a("")
    a("Huawei meaning: when ON, the BSC dynamically lengthens the SDCCH (location update) and initial uplink PS access interval to relieve congestion.")
    a("")
    a("KPI: mainly **CSSR / SDCCH congestion / location-update delay**. It can show up as SDCCH drop or assignment failure. Combined with T3122 below, this is the largest access-policy change in the dump.")
    a("")
    a("### 2. T3122 (`IMMREJWAITINDTIMER`) raised 10 → 15 on 24,175 cells (97.4%)")
    a("")
    a("Huawei default and recommended value is **10**. After an immediate assignment reject, the MS waits T3122 before retrying. A larger value makes access harder. Radio impact (Huawei): *If set large, an MS is difficult to access the network.*")
    a("")
    a("### 3. LTE SAI identity rewritten for SRVCC on ~24,046–24,627 cells (~97–99%)")
    a("")
    a("| Parameter | Typical Pre → Post | Cells |")
    a("|---|---|---:|")
    a("| `LTESAIMCC` | 000 → **470** | 24,046 |")
    a("| `LTESAIMNC` | 000 → **02** | 24,046 |")
    a("| `LTESAILAC` | 1 (or old LAC) → **15000** | 24,624 |")
    a("| `LTESAISAC` | 0 → **15** | 24,627 |")
    a("")
    a("These four values are the SAI used to recognize an **incoming SRVCC handover**. If they do not match what LTE/MME sends, SRVCC is handled as a normal HO (or the reverse). That is a direct CS-continuity / VoLTE→GSM drop mechanism. `SRVCCHOEN` itself changed on only 3 cells (NO→YES), so this SAI rewrite was applied on top of the existing SRVCC-allowed flag.")
    a("")
    a("### 4. VAMOS — Huawei states this increases call drop")
    a("")
    a("- `VAMOSSWITCH` **OFF → ON** on **358 cells**. Huawei: *VAMOS increases capacity by sacrificing quality; quality-related KPIs deteriorate.*")
    a("- About **50 further VAMOS mux/demux parameters** were filled from empty → operating values on those same ~358 cells (first-time VAMOS activation).")
    a("- Load demux thresholds (`VFRLOADREUSETHD`, `VAMOSLOADREUSELOADTHD`) moved on **~17,187 cells (69%)** in ±1 steps (SON-like). Huawei: a smaller demux threshold keeps VAMOS paired longer and **increases call drop rate more significantly**.")
    a("")
    a("### 5. Cluster changes that also hurt quality")
    a("")
    a("| Change | Cells | Why it matters |")
    a("|---|---:|---|")
    a("| Fast-return RSRP `FDDFASTRETURNRSRPTH` 32→22 and `TDDFASTRETURNRSRPTH` 28→22 | 603 | Easier GSM→LTE fast return; Huawei warns of ping-pong. Rec=28. |")
    a("| `SRVCCRAPIDSELMEASOPTSW` OFF→ON | 603 | SRVCC fast-return measurement optimization |")
    a("| `TCHBUSYTHRES` lowered (e.g. 65→20, 80→30, many →0) | 203 | More TCH/H at low load; Huawei: voice quality deteriorates. Rec=60 |")
    a("| `AMRTCHHPRIORLOAD` lowered similarly | 195 | More AMR HR; Rec=55 |")
    a("| `TIGHTBCCHSWITCH` OFF→ON + related HO/assign thresholds | 75 | Aggressive BCCH reuse |")
    a("| `NCC` / `BCC` (BSIC) changed | 636 / 624 | HO failure if neighbors not updated |")
    a("| Energy-save `DYNOPENTRXPOWER` YES→NO | 480 | TRX intelligent shutdown **disabled** (usually helps, not hurts, CDR) |")
    a("")
    a("## Network-wide bulk parameter list (max changes first)")
    a("")
    a("| Rank | MO | Parameter ID | Parameter Name | Cells | % | Pre → Post | KPI | Rec |")
    a("|---:|---|---|---|---:|---:|---|---|---|")
    for i, r in enumerate(bulk, 1):
        a(
            f"| {i} | {r['mo']} | `{r['pid']}` | {r['pname']} | {r['n_changed']:,} | {r['pct']} | {r['transitions'].split(';')[0]} | {r['kpi']} | {r.get('recommended','')} |"
        )
    a("")
    a("## Full changed-parameter list (max changes first)")
    a("")
    a("| Rank | MO | Parameter ID | Parameter Name | Cells Changed | % Common | Change type | Top Pre → Post | KPI |")
    a("|---:|---|---|---|---:|---:|---|---|---|")
    for i, r in enumerate(rows, 1):
        trans = r["transitions"].replace("|", "/")
        a(
            f"| {i} | {r['mo']} | `{r['pid']}` | {r['pname']} | {r['n_changed']:,} | {r['pct']} | {r['change_type']} | {trans} | {r['kpi']} |"
        )
    a("")
    a("## Descriptions of the top bulk changes (from Huawei parameter list)")
    a("")
    for i, r in enumerate(rows[:15], 1):
        a(f"### {i}. {r['mo']} / `{r['pid']}` — {r['pname']}")
        a("")
        a(f"- Cells changed: **{r['n_changed']:,}** ({r['pct']} of {r['n_common']:,})")
        a(f"- Value changes: {r['transitions']}")
        if r.get("default") or r.get("recommended"):
            a(f"- Default: {r.get('default','')} · Recommended: {r.get('recommended','')}")
        if r.get("meaning"):
            a(f"- Meaning: {clip(r['meaning'], 700)}")
        if r.get("impact"):
            a(f"- Radio impact: {clip(r['impact'], 500)}")
        a("")
    a("## MO-level change volume")
    a("")
    a("| MO | Description | Params changed | Total cell-param diffs |")
    a("|---|---|---:|---:|")
    for s in sorted(stats, key=lambda x: -x["n_total_changes"]):
        if s["n_total_changes"] == 0:
            continue
        a(f"| {s['mo']} | {s['mo_desc']} | {s['n_params_changed']} | {s['n_total_changes']:,} |")
    a("")
    a("## What did **not** change")
    a("")
    a("No differences on common cells for handover control/emergency/fast HO (`GCELLHOCTRL`, `GCELLHOEMG`, `GCELLHOFAST`, `GCELLHOFITPEN`), power control (`GCELLPWRBASIC`, `GCELLPWR2`, `GCELLPWR3`), AMR, idle/reselect, GPRS, and most other MOs. That strongly suggests the CDR step is from the bulk campaigns above rather than from a general HO retune.")
    a("")
    a("## Suggested verification order")
    a("")
    a("1. Confirm who enabled **Intelligent Access Congestion Control** and why T3122 was moved off the recommended value 10, with Timer2=100 and threshold=5.")
    a("2. Confirm the **LTE SAI (470/02, LAC 15000, SAC 15)** against the live LTE SRVCC SAI. Mismatch → SRVCC drops.")
    a("3. Correlate TCH CDR by cell with **new VAMOS ON** (358 cells) and with cells whose VAMOS demux load threshold drifted.")
    a("4. Check HO fail after **NCC/BCC** changes (need neighbor BSIC alignment).")
    a("5. Review cells where **TCHBUSYTHRES / AMRTCHHPRIORLOAD** were cut toward 0 (forced HR).")
    a("")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    cat = load_catalog()
    print("catalog", len(cat))
    data = json.loads(JSON_PATH.read_text())
    data = enrich(data, cat)
    JSON_PATH.write_text(json.dumps(data), encoding="utf-8")
    write_csv(data["summary"])
    write_xlsx(data)
    write_md(data)
    print("rows", len(data["summary"]))
    print("wrote", OUT_XLSX, OUT_CSV, OUT_MD)


if __name__ == "__main__":
    main()
